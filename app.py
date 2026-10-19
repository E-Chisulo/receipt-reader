import base64
import hashlib
import io
import json
import os
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

APP_DIR = Path.home() / "ReceiptReader"
WORKBOOK_PATH = Path(os.environ.get("RECEIPT_READER_XLSX", APP_DIR / "receipts.xlsx"))
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://localhost:11434")
MODEL = os.environ.get("OLLAMA_MODEL", "qwen2.5vl:3b")
AUTO_SAVE_SCORE = 0.85
CURRENCY_DEFAULT = "ZMW"

PREDEFINED_CATEGORIES = [
    "Food", "Fuel", "Health", "Internet", "Auto Servicing",
    "Transport", "Printing", "Laundry", "Barber", "Other"
]
CATEGORY_SHEET_HEADERS = ["Date", "Vendor", "Total", "Currency", "Receipt ID"]

RECEIPT_HEADERS = [
    "receipt_id", "captured_at", "merchant", "receipt_date", "currency",
    "subtotal", "tax", "discount", "total", "payment_method", "receipt_number",
    "category", "validation_status", "validation_notes", "image_sha256"
]
ITEM_HEADERS = [
    "receipt_id", "line_number", "item_description", "quantity", "unit_price",
    "line_total", "tax_amount", "category"
]

MEMORY_FILE = APP_DIR / "merchant_memory.json"
APP_DIR.mkdir(parents=True, exist_ok=True)


def load_memory() -> dict:
    if MEMORY_FILE.exists():
        try:
            with open(MEMORY_FILE) as f:
                return json.load(f)
        except Exception:
            pass
    return {"corrections": [], "notes": []}


def save_memory(memory: dict):
    with open(MEMORY_FILE, "w") as f:
        json.dump(memory, f, indent=2)


def add_feedback(feedback: str, merchant: str = None):
    memory = load_memory()
    entry = {
        "date": datetime.now().strftime("%Y-%m-%d"),
        "feedback": feedback,
        "merchant": merchant,
    }
    memory["corrections"].append(entry)
    save_memory(memory)


def memory_prompt_section() -> str:
    memory = load_memory()
    corrections = memory.get("corrections", [])
    if not corrections:
        return ""
    recent = corrections[-10:]
    lines = ["Past corrections and notes from the user (apply these when relevant):"]
    for c in recent:
        merchant_ctx = f" (receipt from: {c['merchant']})" if c.get("merchant") else ""
        lines.append(f"- {c['feedback']}{merchant_ctx}")
    return "\n".join(lines)


def _style_header(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B1220")


def ensure_category_sheet(wb, category: str):
    if category not in wb.sheetnames:
        ws = wb.create_sheet(category)
        ws.append(CATEGORY_SHEET_HEADERS)
        _style_header(ws)
    return wb[category]


def ensure_workbook(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Receipts"
        ws.append(RECEIPT_HEADERS)
        _style_header(ws)
        ws_items = wb.create_sheet("Items")
        ws_items.append(ITEM_HEADERS)
        _style_header(ws_items)
        for cat in PREDEFINED_CATEGORIES:
            cs = wb.create_sheet(cat)
            cs.append(CATEGORY_SHEET_HEADERS)
            _style_header(cs)
        wb.save(path)
    else:
        wb = load_workbook(path)
        changed = False
        if "Receipts" not in wb.sheetnames:
            ws = wb.create_sheet("Receipts", 0)
            ws.append(RECEIPT_HEADERS)
            _style_header(ws)
            changed = True
        if "Items" not in wb.sheetnames:
            ws = wb.create_sheet("Items")
            ws.append(ITEM_HEADERS)
            _style_header(ws)
            changed = True
        for cat in PREDEFINED_CATEGORIES:
            if cat not in wb.sheetnames:
                cs = wb.create_sheet(cat)
                cs.append(CATEGORY_SHEET_HEADERS)
                _style_header(cs)
                changed = True
        if changed:
            wb.save(path)


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def image_to_jpeg_bytes(image_bytes: bytes) -> bytes:
    image = Image.open(io.BytesIO(image_bytes))
    image = ImageOps.exif_transpose(image).convert("RGB")
    # Downsize very large images to reduce local inference time while retaining legibility.
    image.thumbnail((1400, 1400))
    out = io.BytesIO()
    image.save(out, format="JPEG", quality=92)
    return out.getvalue()


def ollama_extract(jpeg_bytes: bytes) -> dict:
    categories_list = ", ".join(PREDEFINED_CATEGORIES)
    memory_section = memory_prompt_section()
    prompt = f"""
Read this receipt carefully. Return ONLY a valid JSON object, no markdown.
Do not guess unreadable values: use null for unknown values and include them in
"uncertainties". Amounts must be numbers, not strings. Use decimal numbers.
Return this exact schema:
{{
  "merchant": string|null,
  "receipt_date": "YYYY-MM-DD"|null,
  "currency": string|null,
  "subtotal": number|null,
  "tax": number|null,
  "discount": number|null,
  "total": number|null,
  "payment_method": string|null,
  "receipt_number": string|null,
  "category": string,
  "items": [
    {{
      "description": string,
      "quantity": number|null,
      "unit_price": number|null,
      "line_total": number|null,
      "tax_amount": number|null,
      "category": string|null
    }}
  ],
  "uncertainties": [string],
  "model_confidence": number
}}
Rules:
- Capture every visible purchased line item, including quantity and amount.
- Distinguish unit price from line total. If the receipt only gives a line total,
  leave unit_price null unless it is clear.
- Do not treat change, tendered cash, card balance, or loyalty points as the total.
- If a printed tax is absent, use null, not zero.
- model_confidence is your own estimate from 0 to 1, but it is not independently verified.
- MERCHANT IDENTIFICATION: The merchant is the business selling the goods or services.
  Bank or payment provider names printed on receipts (FNB, ABSA, Zanaco, Standard Bank,
  Stanbic, Visa, Mastercard, MasterPass, PayPak) are payment infrastructure — NOT the merchant.
  Ignore bank logos and card scheme names. Look for the actual business name, which is
  usually printed near the top separately from the bank branding.
- For "category", classify the entire receipt into one of: {categories_list}.
  Use the items to infer category even if the merchant name is missing or unclear.
  Examples: groceries/restaurants/cafes → Food; petrol stations/diesel → Fuel;
  pharmacy/clinic/hospital/medicine → Health; mobile data/fibre/wifi → Internet;
  car wash/tyre/mechanic → Auto Servicing; taxi/ride/bus/flight/ticket → Transport;
  photocopy/document printing → Printing; laundry/dry clean → Laundry;
  haircut/grooming → Barber. If none fit clearly, use Other.
  For item-level "category", use the same list to classify each line item individually.
{memory_section}
"""
    payload = {
        "model": MODEL,
        "stream": False,
        "format": "json",
        "messages": [{
            "role": "user",
            "content": prompt,
            "images": [base64.b64encode(jpeg_bytes).decode("ascii")]
        }],
        "options": {"temperature": 0, "num_ctx": 8192}
    }
    try:
        response = requests.post(f"{OLLAMA_URL}/api/chat", json=payload, timeout=240)
        response.raise_for_status()
        body = response.json()
        content = body.get("message", {}).get("content", "")
        result = json.loads(content)
        if not isinstance(result, dict):
            raise ValueError("Model response was not a JSON object.")
        return result
    except requests.exceptions.ConnectionError as exc:
        raise RuntimeError(
            f"Cannot reach Ollama at {OLLAMA_URL}. Open Ollama and confirm the model "
            f"'{MODEL}' is installed."
        ) from exc
    except requests.exceptions.Timeout as exc:
        raise RuntimeError("The local model took too long. Try again or use a smaller model.") from exc
    except (requests.HTTPError, ValueError, json.JSONDecodeError) as exc:
        detail = ""
        try:
            detail = response.text[:500]
        except Exception:
            pass
        raise RuntimeError(f"Could not parse a valid receipt result: {exc}. {detail}") from exc


def num(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def money_close(a, b, tolerance=0.03):
    return a is not None and b is not None and abs(a - b) <= tolerance


def validate_receipt(data: dict):
    notes = []
    score = num(data.get("model_confidence"))
    if score is None:
        score = 0
    score = max(0.0, min(1.0, score))
    items = data.get("items")
    if not isinstance(items, list) or len(items) == 0:
        notes.append("No line items extracted.")
        items = []
    if not data.get("merchant"):
        notes.append("Merchant is missing.")
    if not data.get("receipt_date"):
        notes.append("Receipt date is missing.")
    total = num(data.get("total"))
    if total is None or total <= 0:
        notes.append("Receipt total is missing or invalid.")

    # Validate item arithmetic where the necessary values exist.
    known_line_totals = []
    for idx, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            notes.append(f"Line {idx} is not a valid item object.")
            continue
        description = item.get("description")
        if not description:
            notes.append(f"Line {idx} has no description.")
        qty = num(item.get("quantity"))
        unit = num(item.get("unit_price"))
        line_total = num(item.get("line_total"))
        if qty is not None and qty <= 0:
            notes.append(f"Line {idx} has a non-positive quantity.")
        if unit is not None and unit < 0:
            notes.append(f"Line {idx} has a negative unit price.")
        if line_total is not None and line_total < 0:
            notes.append(f"Line {idx} has a negative line total.")
        if qty is not None and unit is not None and line_total is not None:
            if not money_close(qty * unit, line_total, 0.05):
                notes.append(f"Line {idx} quantity × unit price does not match line total.")
        if line_total is not None:
            known_line_totals.append(line_total)

    subtotal = num(data.get("subtotal"))
    tax = num(data.get("tax"))
    discount = num(data.get("discount"))
    if subtotal is not None and subtotal < 0:
        notes.append("Subtotal is negative.")
    if tax is not None and tax < 0:
        notes.append("Tax is negative.")
    if total is not None and subtotal is not None and tax is not None:
        # Common receipt equation; allow for discounts when shown.
        expected = subtotal + tax - (discount or 0)
        if not money_close(expected, total, 0.10):
            notes.append("Subtotal + tax − discount does not match total.")
    if subtotal is not None and len(known_line_totals) == len(items) and items:
        if not money_close(sum(known_line_totals), subtotal, 0.10):
            notes.append("Sum of line totals does not match subtotal.")

    uncertainties = data.get("uncertainties")
    if isinstance(uncertainties, list) and uncertainties:
        notes.append("Model reported uncertainties: " + "; ".join(str(x) for x in uncertainties[:5]))

    # Conservative automatic save: model confidence plus required data and consistency checks.
    hard_missing = (
        not data.get("merchant") or not data.get("receipt_date") or
        total is None or total <= 0 or not items
    )
    save_automatically = score >= AUTO_SAVE_SCORE and not hard_missing and not notes
    return save_automatically, score, notes


def find_content_duplicate(path: Path, merchant: str, date: str, total: float) -> str | None:
    """Return receipt_id if a receipt with same merchant+date+total already exists."""
    if not path.exists() or not merchant or not date or total is None:
        return None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Receipts"]
        headers = [c.value for c in ws[1]]
        def col(name):
            return headers.index(name) if name in headers else None
        m_col, d_col, t_col, id_col = col("merchant"), col("receipt_date"), col("total"), col("receipt_id")
        if None in (m_col, d_col, t_col, id_col):
            return None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if (str(row[m_col] or "").strip().lower() == merchant.strip().lower()
                    and str(row[d_col] or "") == str(date)
                    and row[t_col] is not None
                    and abs(float(row[t_col]) - float(total)) < 0.01):
                return str(row[id_col])
        wb.close()
    except Exception:
        pass
    return None


def existing_hashes(path: Path):
    ensure_workbook(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Receipts"]
    headers = [c.value for c in ws[1]]
    idx = headers.index("image_sha256") + 1 if "image_sha256" in headers else None
    hashes = set()
    if idx:
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx, values_only=True):
            if row[0]:
                hashes.add(row[0])
    wb.close()
    return hashes


def append_to_excel(data: dict, path: Path, digest: str, status: str, notes: list):
    ensure_workbook(path)
    wb = load_workbook(path)
    ws = wb["Receipts"]
    item_ws = wb["Items"]
    receipt_id = datetime.now().strftime("R%Y%m%d%H%M%S%f")
    captured_at = datetime.now().astimezone().isoformat(timespec="seconds")
    category = data.get("category") or "Other"
    row = {
        "receipt_id": receipt_id,
        "captured_at": captured_at,
        "merchant": data.get("merchant"),
        "receipt_date": data.get("receipt_date"),
        "currency": data.get("currency") or CURRENCY_DEFAULT,
        "subtotal": num(data.get("subtotal")),
        "tax": num(data.get("tax")),
        "discount": num(data.get("discount")),
        "total": num(data.get("total")),
        "payment_method": data.get("payment_method"),
        "receipt_number": data.get("receipt_number"),
        "category": category,
        "validation_status": status,
        "validation_notes": "; ".join(notes),
        "image_sha256": digest,
    }
    ws.append([row.get(h) for h in RECEIPT_HEADERS])

    # Write to category sheet (auto-create if new category)
    cat_ws = ensure_category_sheet(wb, category)
    cat_ws.append([
        data.get("receipt_date"),
        data.get("merchant") or "Unknown",
        num(data.get("total")),
        data.get("currency") or CURRENCY_DEFAULT,
        receipt_id,
    ])
    for i, item in enumerate(data.get("items") or [], start=1):
        if not isinstance(item, dict):
            continue
        item_row = {
            "receipt_id": receipt_id,
            "line_number": i,
            "item_description": item.get("description"),
            "quantity": num(item.get("quantity")),
            "unit_price": num(item.get("unit_price")),
            "line_total": num(item.get("line_total")),
            "tax_amount": num(item.get("tax_amount")),
            "category": item.get("category"),
        }
        item_ws.append([item_row.get(h) for h in ITEM_HEADERS])
    for sheet in (ws, item_ws):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col_idx in range(1, sheet.max_column + 1):
            max_len = 0
            for cell in sheet[get_column_letter(col_idx)]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 38)
    wb.save(path)
    return receipt_id


def to_display_frame(data):
    items = data.get("items") or []
    return pd.DataFrame([{
        "Item": x.get("description"),
        "Quantity": x.get("quantity"),
        "Unit price": x.get("unit_price"),
        "Line total": x.get("line_total"),
        "Tax amount": x.get("tax_amount"),
        "Category": x.get("category"),
    } for x in items if isinstance(x, dict)])


st.set_page_config(page_title="Receipt Reader", page_icon="🧾", layout="wide")
st.markdown("""
<style>
/* Camera widget container */
[data-testid="stCameraInput"] > div {
    width: 100% !important;
}
/* The inner box that holds the video */
[data-testid="stCameraInput"] > div > div {
    width: 100% !important;
    aspect-ratio: 3 / 4 !important;
    height: auto !important;
    min-height: 60vh !important;
}
/* The video element itself */
[data-testid="stCameraInput"] video {
    width: 100% !important;
    height: 100% !important;
    aspect-ratio: 3 / 4 !important;
    object-fit: cover !important;
    border-radius: 8px;
    min-height: 60vh !important;
}
/* Canvas used for capture */
[data-testid="stCameraInput"] canvas {
    width: 100% !important;
    aspect-ratio: 3 / 4 !important;
    min-height: 60vh !important;
}
</style>
""", unsafe_allow_html=True)
st.title("🧾 Receipt Reader")
st.caption("Local receipt extraction • Ollama vision model • Excel workbook on this Mac")

with st.sidebar:
    st.subheader("Settings")
    workbook_text = st.text_input("Excel workbook path", str(WORKBOOK_PATH))
    model_text = st.text_input("Ollama model", MODEL)
    st.caption("Default model: qwen2.5vl:3b")
    st.caption("Images are processed locally by Ollama. No cloud AI API is used.")
    st.divider()
    auto_mode = st.toggle("⚡ Auto Mode", value=st.session_state.get("auto_mode", False),
                          help="Reads receipt the moment a photo is taken — no button needed.")
    st.session_state["auto_mode"] = auto_mode
    if auto_mode:
        st.caption("Photo → AI reads → you approve (or it saves automatically).")
    st.divider()
    st.write("**Automatic-save policy**")
    st.write("Saves only when the model score is ≥ 0.85 and required fields and arithmetic checks pass.")
    st.caption("Model confidence is an estimate, not a calibrated probability.")

WORKBOOK_PATH = Path(workbook_text).expanduser()
MODEL = model_text.strip() or MODEL
ensure_workbook(WORKBOOK_PATH)

st.write("Position the receipt clearly in front of the camera, then capture one image. Use bright, even lighting and keep the full receipt in frame.")
tab_upload, tab_cam = st.tabs(["📷 Take Photo", "🖥️ Mac Camera"])
with tab_upload:
    st.caption("Tap **Browse** → **Take Photo** for best focus and quality using your iPhone camera.")
    uploaded = st.file_uploader("Receipt image", type=["jpg", "jpeg", "png", "webp", "heic"], label_visibility="collapsed")
    image_file = None
    from_upload = False
    if uploaded is not None:
        image_file = uploaded
        from_upload = True
with tab_cam:
    cam_img = st.camera_input("Capture receipt with your Mac camera")
    if cam_img is not None:
        image_file = cam_img
        from_upload = False

if image_file is not None:
    original_bytes = image_file.getvalue()
    digest = sha256(original_bytes)
    # New photo arrived — reset processed digest so auto mode fires on the new image
    if digest != st.session_state.get("last_processed_digest"):
        if "last_receipt" not in st.session_state:
            st.session_state.pop("last_processed_digest", None)
    if from_upload:
        st.image(original_bytes, caption="Uploaded receipt", use_container_width=True)
    if digest in existing_hashes(WORKBOOK_PATH):
        st.warning("This exact image is already in the workbook. It will not be added again.")
    else:
        already_extracted = (
            st.session_state.get("last_receipt", {}).get("digest") == digest or
            st.session_state.get("last_processed_digest") == digest
        )
        trigger = (
            not already_extracted and auto_mode
        ) or (
            not auto_mode and st.button("Read receipt with local AI", type="primary")
        )
        if auto_mode and not already_extracted:
            st.info("⚡ Auto mode — reading receipt…")
        if trigger:
            with st.spinner(f"Reading receipt locally with {MODEL}…"):
                try:
                    t0 = datetime.now()
                    jpeg = image_to_jpeg_bytes(original_bytes)
                    extracted = ollama_extract(jpeg)
                    elapsed = (datetime.now() - t0).total_seconds()
                    auto_save, score, notes = validate_receipt(extracted)
                    dup_id = find_content_duplicate(
                        WORKBOOK_PATH,
                        extracted.get("merchant"),
                        extracted.get("receipt_date"),
                        num(extracted.get("total"))
                    )
                    st.session_state["last_processed_digest"] = digest
                    st.session_state["last_receipt"] = {
                        "data": extracted, "digest": digest, "auto_save": auto_save,
                        "score": score, "notes": notes, "workbook": str(WORKBOOK_PATH),
                        "elapsed": elapsed, "dup_id": dup_id
                    }
                    st.rerun()
                except Exception as e:
                    st.error(str(e))

if "last_receipt" in st.session_state:
    result = st.session_state["last_receipt"]
    data = result["data"]
    if result["workbook"] != str(WORKBOOK_PATH):
        st.info("The workbook path changed since this receipt was read. Choose where to save below.")

    dup_id = result.get("dup_id")
    if dup_id and not result.get("saved"):
        st.error(f"⚠️ A receipt with the same merchant, date and total is already saved (ID: {dup_id}). Check before saving again.")

    st.subheader("Extracted receipt")
    elapsed = result.get("elapsed")
    timing = f" · inference {elapsed:.1f}s" if elapsed else ""
    st.caption(f"Model confidence estimate: {result['score']:.0%}{timing}")

    receipt_date = data.get("receipt_date") or "unknown"
    merchant = data.get("merchant") or "unknown vendor"

    def parse_date(d: str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(d, fmt)
            except Exception:
                pass
        return None

    def spoken_date(d: str) -> str:
        dt = parse_date(d)
        if not dt:
            return d
        day = dt.day
        suffix = "th" if 11 <= day <= 13 else {1:"st",2:"nd",3:"rd"}.get(day % 10, "th")
        return dt.strftime(f"{day}{suffix} %B %Y")

    def display_date(d: str) -> str:
        dt = parse_date(d)
        if not dt:
            return d
        return dt.strftime(f"%-d %b, %Y")

    speak_text = f"Receipt date: {spoken_date(receipt_date)}. Merchant: {merchant}."
    if not result.get("spoken"):
        import streamlit.components.v1 as components
        components.html(f"""
        <script>
        var u = new SpeechSynthesisUtterance({json.dumps(speak_text)});
        u.rate = 0.95;
        window.speechSynthesis.cancel();
        window.speechSynthesis.speak(u);
        </script>
        """, height=0)
        result["spoken"] = True
        st.session_state["last_receipt"] = result
    currency = data.get("currency") or CURRENCY_DEFAULT
    total = num(data.get("total"))
    total_str = f"{currency} {total:,.2f}" if total is not None else "—"

    st.markdown(f"""
<div style="display:flex;gap:12px;margin:12px 0;">
  <div style="flex:1;background:#1e3a5f;border-radius:10px;padding:16px 12px;text-align:center;">
    <div style="color:#aac4e0;font-size:0.75rem;margin-bottom:4px;">VENDOR</div>
    <div style="color:#ffffff;font-size:1.1rem;font-weight:700;word-break:break-word;">{merchant}</div>
  </div>
  <div style="flex:1;background:#1a4731;border-radius:10px;padding:16px 12px;text-align:center;">
    <div style="color:#8ecfaf;font-size:0.75rem;margin-bottom:4px;">AMOUNT</div>
    <div style="color:#4ade80;font-size:1.2rem;font-weight:700;">{total_str}</div>
  </div>
  <div style="flex:1;background:#3b2a1a;border-radius:10px;padding:16px 12px;text-align:center;">
    <div style="color:#d4a96a;font-size:0.75rem;margin-bottom:4px;">DATE</div>
    <div style="color:#fbbf24;font-size:0.95rem;font-weight:700;">{display_date(receipt_date)}</div>
  </div>
</div>
""", unsafe_allow_html=True)

    if result["auto_save"]:
        st.success("Validation passed. This receipt is eligible for automatic saving.")
        if not result.get("saved"):
            try:
                rid = append_to_excel(data, WORKBOOK_PATH, result["digest"], "auto-saved", [])
                result["saved"] = True
                result["receipt_id"] = rid
                st.session_state["last_receipt"] = result
                cat = data.get("category") or "Other"
                st.toast(f"✓ Saved · {merchant} · {total_str} · {cat}", icon="✅")
                import time; time.sleep(1.2)
                del st.session_state["last_receipt"]
                st.rerun()
            except Exception as e:
                st.error(f"Could not write Excel file: {e}")
    else:
        st.warning("Held for review. The app will not automatically save this receipt.")
        uncertainty_notes = [n for n in result["notes"] if n.startswith("Model reported uncertainties")]
        validation_notes = [n for n in result["notes"] if not n.startswith("Model reported uncertainties")]
        if result["score"] < AUTO_SAVE_SCORE:
            st.write(f"- Confidence estimate ({result['score']:.0%}) is below the {AUTO_SAVE_SCORE:.0%} threshold.")
        if validation_notes:
            st.write("**Validation issues:**")
            for note in validation_notes:
                st.write(f"- {note}")
        if uncertainty_notes:
            st.write("**Model uncertainties:**")
            for n in uncertainty_notes:
                raw = n.replace("Model reported uncertainties: ", "")
                for item in raw.split("; "):
                    if item.strip():
                        st.write(f"- {item.strip()}")
    col_left, col_right = st.columns(2)
    with col_left:
        edited_merchant = st.text_input("Merchant", value=str(data.get("merchant") or ""))
        edited_date = st.text_input("Receipt date (YYYY-MM-DD)", value=str(data.get("receipt_date") or ""))
        edited_currency = st.text_input("Currency", value=str(data.get("currency") or CURRENCY_DEFAULT))
    with col_right:
        inferred_cat = data.get("category") or "Other"
        all_cats = PREDEFINED_CATEGORIES + (
            [inferred_cat] if inferred_cat not in PREDEFINED_CATEGORIES else []
        )
        default_idx = all_cats.index(inferred_cat) if inferred_cat in all_cats else 0
        edited_category = st.selectbox(
            "Category (AI inferred — correct if needed)",
            options=all_cats,
            index=default_idx,
        )
        st.caption(f"AI classified this as: **{inferred_cat}**")
    col1, col2, col3 = st.columns(3)
    with col1:
        edited_subtotal = st.number_input("Subtotal", min_value=0.0, value=float(num(data.get("subtotal")) or 0.0), step=1.0)
    with col2:
        edited_tax = st.number_input("Tax", min_value=0.0, value=float(num(data.get("tax")) or 0.0), step=1.0)
    with col3:
        edited_total = st.number_input("Total", min_value=0.0, value=float(num(data.get("total")) or 0.0), step=1.0)
    st.write("**Line items**")
    items_df = to_display_frame(data)
    edited_items = st.data_editor(items_df, num_rows="dynamic", use_container_width=True, key="items_editor")
    with st.expander("Raw extracted JSON"):
        st.json(data)

    with st.expander("💬 Give feedback to improve future extractions"):
        st.caption("Corrections you type here are saved and injected into every future extraction prompt.")
        fb_text = st.text_area(
            "Feedback",
            placeholder='e.g. "FNB is the bank, not the merchant — look for the business name above it"\n"Med One Pharmacy Woodlands is the correct merchant for FNB receipts from Woodlands"\n"Receipts showing K prefix are in ZMW"',
            key="feedback_input",
            height=100,
        )
        if st.button("Save feedback", key="save_feedback"):
            if fb_text.strip():
                add_feedback(fb_text.strip(), merchant=data.get("merchant"))
                st.success("Saved — will apply to the next extraction.")
            else:
                st.warning("Type something first.")

        memory = load_memory()
        if memory["corrections"]:
            st.write(f"**{len(memory['corrections'])} correction(s) saved:**")
            for c in reversed(memory["corrections"][-5:]):
                st.write(f"- *{c['date']}*: {c['feedback']}")

    if not result["auto_save"]:
        st.markdown("""
<style>
#approve-bar { position:fixed; bottom:0; left:0; right:0; z-index:9999;
  padding:12px 16px; background:#0e1117;
  border-top:1px solid #333; box-shadow:0 -2px 12px rgba(0,0,0,0.4); }
#approve-bar button {
  width:100%; padding:14px; background:#ff4b4b; color:#fff;
  border:none; border-radius:8px; font-size:1rem; font-weight:700; cursor:pointer; }
#approve-bar button:active { background:#cc3a3a; }
.approve-spacer { height: 80px; }
</style>
<div class="approve-spacer"></div>
<div id="approve-bar">
  <button id="approve-float" onclick="
    var self = this;
    self.disabled = true;
    self.innerText = 'Saving…';
    var found = false;
    var allBtns = Array.from(window.parent.document.querySelectorAll('button'));
    if (!allBtns.length) allBtns = Array.from(document.querySelectorAll('button'));
    for (var i = 0; i < allBtns.length; i++) {
      if (allBtns[i].innerText.includes('Approve')) {
        allBtns[i].dispatchEvent(new MouseEvent('click', {bubbles:true, cancelable:true, view:window}));
        found = true;
        break;
      }
    }
    if (!found) { self.disabled = false; self.innerText = '✓ Approve and Save to Excel'; }
  ">✓ Approve and Save to Excel</button>
</div>
""", unsafe_allow_html=True)
    if not result["auto_save"] and st.button("Approve and save to Excel", type="primary", key="approve_hidden") and not result.get("saved"):
        corrected = dict(data)
        corrected["merchant"] = edited_merchant.strip() or None
        corrected["receipt_date"] = edited_date.strip() or None
        corrected["currency"] = edited_currency.strip() or CURRENCY_DEFAULT
        corrected["category"] = edited_category
        corrected["subtotal"] = edited_subtotal
        corrected["tax"] = edited_tax
        corrected["total"] = edited_total
        corrected_items = []
        for _, row in edited_items.iterrows():
            if pd.isna(row.get("Item")) or not str(row.get("Item")).strip():
                continue
            corrected_items.append({
                "description": str(row.get("Item")),
                "quantity": num(row.get("Quantity")),
                "unit_price": num(row.get("Unit price")),
                "line_total": num(row.get("Line total")),
                "tax_amount": num(row.get("Tax amount")),
                "category": None if pd.isna(row.get("Category")) else str(row.get("Category")),
            })
        corrected["items"] = corrected_items
        rid = append_to_excel(corrected, WORKBOOK_PATH, result["digest"], "manually-approved", ["User reviewed and approved"])
        saved_merchant = corrected.get("merchant") or "Receipt"
        saved_total = f"{corrected.get('currency', CURRENCY_DEFAULT)} {corrected.get('total', 0):,.2f}"
        st.toast(f"✓ Saved · {saved_merchant} · {saved_total}", icon="✅")
        import time; time.sleep(1.2)
        del st.session_state["last_receipt"]
        st.rerun()

st.divider()
st.subheader("Workbook")
st.write(f"Excel file: `{WORKBOOK_PATH}`")
if WORKBOOK_PATH.exists():
    try:
        receipts_df = pd.read_excel(WORKBOOK_PATH, sheet_name="Receipts")
        items_df = pd.read_excel(WORKBOOK_PATH, sheet_name="Items")
        col1, col2 = st.columns(2)
        col1.metric("Receipts recorded", max(0, len(receipts_df)))
        col2.metric("Line items recorded", max(0, len(items_df)))

        if "category" in receipts_df.columns and len(receipts_df) > 0:
            with st.expander("Spending by category"):
                cat_summary = (
                    receipts_df.groupby("category")["total"]
                    .agg(count="count", total="sum")
                    .reset_index()
                    .sort_values("total", ascending=False)
                )
                cat_summary.columns = ["Category", "Receipts", "Total (ZMW)"]
                st.dataframe(cat_summary, use_container_width=True, hide_index=True)

        with st.expander("Recent receipts"):
            st.dataframe(receipts_df.tail(10), use_container_width=True)
        with st.expander("Recent line items"):
            st.dataframe(items_df.tail(20), use_container_width=True)
    except Exception as e:
        st.info(f"Workbook preview unavailable: {e}")
